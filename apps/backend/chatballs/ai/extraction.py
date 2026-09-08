"""Извлечение текста из файловых вложений знаний (ADR-CHATBALLS-0023).

Поддерживаются текстовые форматы (md, txt и любой text/*), PDF (pypdf) и DOCX
(python-docx). Остальные форматы хранятся без индексации — extract_text вернёт
пустую строку. Ошибки разбора не фатальны: вложение сохраняется, текст пустой.
"""

from __future__ import annotations

import io
import logging
from pathlib import PurePosixPath

logger = logging.getLogger(__name__)

_TEXT_EXTENSIONS = {".md", ".markdown", ".txt", ".csv", ".json", ".yaml", ".yml"}


def _decode(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(filter(None, (page.extract_text() for page in reader.pages)))


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(*, filename: str, content_type: str, data: bytes) -> str:
    suffix = PurePosixPath(filename.lower()).suffix
    try:
        if suffix in _TEXT_EXTENSIONS or content_type.startswith("text/"):
            return _decode(data).strip()
        if suffix == ".pdf" or content_type == "application/pdf":
            return _extract_pdf(data).strip()
        if suffix == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return _extract_docx(data).strip()
    except Exception:  # noqa: BLE001 — формат-специфичные ошибки разбора не фатальны
        logger.warning("Text extraction failed for %s (%s)", filename, content_type, exc_info=True)
    return ""
